"""Junta as codificações independentes (amostra01_h1.xlsx e amostra01_h2.xlsx) na planilha oficial
02-triagem/validacao/ta_v1/amostra01_cega.xlsx, casando por id_rs. Não toca em nenhuma outra coluna.
Uso: python3 08-revisao-humana/P006_P007_validacao/juntar_h1_h2.py"""
import openpyxl
from pathlib import Path
raiz = Path(__file__).resolve().parents[2]
aqui = Path(__file__).resolve().parent
oficial = raiz / "02-triagem/validacao/ta_v1/amostra01_cega.xlsx"

def ler(arq, cols):
    ws = openpyxl.load_workbook(arq)["codificacao"]
    hdr = [c.value for c in ws[1]]
    return {r[hdr.index("id_rs")]: {c: r[hdr.index(c)] for c in cols}
            for r in ws.iter_rows(min_row=2, values_only=True) if r[hdr.index("id_rs")]}

h1 = ler(aqui / "amostra01_h1.xlsx", ["decisao_h1", "criterio_h1"])
h2 = ler(aqui / "amostra01_h2.xlsx", ["decisao_h2", "criterio_h2"])
wb = openpyxl.load_workbook(oficial); ws = wb["codificacao"]
hdr = [c.value for c in ws[1]]
n = 0
for row in ws.iter_rows(min_row=2):
    id_rs = row[hdr.index("id_rs")].value
    for fonte in (h1, h2):
        for col, val in fonte.get(id_rs, {}).items():
            if val not in (None, ""):
                row[hdr.index(col)].value = val; n += 1
wb.save(oficial)
print(f"{n} células copiadas para {oficial.relative_to(raiz)}; preencha decisao_consenso nas discordâncias e rode validar calcular")
